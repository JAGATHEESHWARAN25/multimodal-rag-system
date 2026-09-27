import os
import docx
from pptx import Presentation
from pptx.util import Inches
import numpy as np
import cv2

OUTPUT_DIR = "demo_dataset"

def create_docx():
    doc = docx.Document()
    doc.add_heading("System Deployment Specifications", 0)
    
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph("The National Surveillance Architecture (NSA) requires highly secure, on-premise infrastructure. Deployed at NTRO headquarters on 2026-11-15.")
    
    doc.add_heading("2. Server Requirements", level=1)
    doc.add_paragraph("All servers must support 16GB-RAM minimum and execute the Florence-2 Vision layer offline.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Specification"
    table.cell(1, 0).text = "Vision Model"
    table.cell(1, 1).text = "microsoft/Florence-2-base"
    
    doc.add_paragraph("Classification: CONFIDENTIAL")
    
    doc.save(os.path.join(OUTPUT_DIR, "System_Deployment_Specifications.docx"))

def create_pptx():
    prs = Presentation()
    
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes
    
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Network Architecture Presentation"
    
    tf = body_shape.text_frame
    tf.text = "Overview of the National Surveillance Architecture"
    
    p = tf.add_paragraph()
    p.text = "- 100% Offline Capable"
    p = tf.add_paragraph()
    p.text = "- Integrated with NTRO systems"
    
    # Add a simple table
    rows = cols = 2
    left = top = Inches(2)
    width = Inches(6)
    height = Inches(1.5)
    
    table_shape = shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.cell(0, 0).text = "Node"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Gateway"
    table.cell(1, 1).text = "Filter"
    
    # Add speaker notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Remember to emphasize that the system runs on CPU-first."
    
    prs.save(os.path.join(OUTPUT_DIR, "Network_Architecture.pptx"))

def create_image():
    # Create a synthetic diagram image
    img = np.ones((600, 800, 3), dtype=np.uint8) * 240
    
    # Draw a box for "API Gateway"
    cv2.rectangle(img, (100, 100), (300, 200), (0, 0, 0), 2)
    cv2.putText(img, "API Gateway", (120, 160), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)
    
    # Draw a box for "Local LLM"
    cv2.rectangle(img, (500, 100), (700, 200), (0, 0, 0), 2)
    cv2.putText(img, "Local LLM", (540, 160), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)
    
    # Draw arrow connecting them
    cv2.arrowedLine(img, (300, 150), (500, 150), (255, 0, 0), 3)
    
    # Draw a box for "ChromaDB"
    cv2.rectangle(img, (300, 400), (500, 500), (0, 0, 0), 2)
    cv2.putText(img, "ChromaDB", (340, 460), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)
    
    cv2.imwrite(os.path.join(OUTPUT_DIR, "System_Architecture_Diagram.png"), img)

def create_pdf():
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt="National Surveillance Architecture", ln=1, align='C')
        pdf.cell(200, 10, txt="Document ID: NTRO-SEC-991", ln=2)
        pdf.multi_cell(0, 10, txt="The National Surveillance Architecture is a secure, isolated platform designed for government deployments. It connects securely to the API Gateway and utilizes the Local LLM.")
        pdf.output(os.path.join(OUTPUT_DIR, "National_Surveillance_Architecture.pdf"))
    except ImportError:
        pass

if __name__ == "__main__":
    if not os.path.exists(OUTPUT_DIR):
        os.makedirs(OUTPUT_DIR)
    
    create_docx()
    create_pptx()
    create_image()
    create_pdf()
    
    print(f"Dataset generated successfully in {OUTPUT_DIR}")

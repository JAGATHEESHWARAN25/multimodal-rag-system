import io
import pytest
import openpyxl
from app.core.ingestion.adapters.xlsx_adapter import XLSXAdapter
from app.core.pipeline.orchestrator import PipelineOrchestrator

def test_xlsx_adapter_multi_sheet():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Servers"
    ws1.append(["Host", "IP"])
    ws1.append(["AUTH-SRV-01", "10.0.0.15"])
    
    ws2 = wb.create_sheet("Network")
    ws2.append(["Device", "Zone"])
    ws2.append(["RADAR-01", "Alpha"])
    
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()
    
    adapter = XLSXAdapter()
    cdo = adapter.parse(xlsx_bytes, {"filename": "infra.xlsx"})
    
    assert cdo.modality == "XLSX"
    assert len(cdo.pages) == 2
    assert cdo.pages[0]["metadata"]["sheet_name"] == "Servers"
    assert cdo.pages[1]["metadata"]["sheet_name"] == "Network"

def test_xlsx_formula_safety():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Value1", "Value2", "SumFormula"])
    ws.append([10, 20, "=SUM(A2:B2)"])
    
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()
    
    adapter = XLSXAdapter()
    cdo = adapter.parse(xlsx_bytes, {"filename": "formula.xlsx"})
    table = cdo.pages[0]["elements"][0]
    
    # Formula cell value string must be treated strictly as data, never executed
    cell_val = table["rows"][0]["cells"][2]["text"]
    assert cell_val in ["=SUM(A2:B2)", "None", "30", ""]

def test_xlsx_orchestration():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ServerName", "IPAddress"])
    ws.append(["AUTH-SRV-01", "10.0.0.15"])
    buf = io.BytesIO()
    wb.save(buf)
    
    adapter = XLSXAdapter()
    cdo = adapter.parse(buf.getvalue(), {"filename": "test.xlsx"})
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    assert schema.source_file == "test.xlsx"
    assert any("cell_coordinate" in c.get("metadata", {}) for c in chunks)

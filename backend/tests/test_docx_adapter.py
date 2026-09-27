import pytest
import io
import docx
from docx.enum.style import WD_STYLE_TYPE
from fastapi.testclient import TestClient

from app.main import app
from app.core.ingestion.adapters.docx_adapter import DOCXAdapter
from app.core.ingestion.engine import IngestionEngine
from app.core.pipeline.orchestrator import PipelineOrchestrator
from app.core.ingestion.detector import CorruptedFileError

client = TestClient(app)

def create_synthetic_docx(test_type="basic") -> bytes:
    doc = docx.Document()
    
    if test_type == "basic":
        doc.add_heading('Title', 0)
        doc.add_paragraph('Paragraph 1')
        doc.add_paragraph('Paragraph 2')
        
    elif test_type == "headings":
        doc.add_heading('Heading 1', level=1)
        doc.add_paragraph('Paragraph under H1')
        doc.add_heading('Heading 2', level=2)
        doc.add_paragraph('Paragraph under H2')
        doc.add_heading('Heading 2 (Another)', level=2)
        doc.add_paragraph('Paragraph under another H2')
        
    elif test_type == "tables":
        doc.add_paragraph('Here is a table:')
        table = doc.add_table(rows=3, cols=4)
        for row in range(3):
            for col in range(4):
                table.cell(row, col).text = f"R{row}C{col}"
                
    elif test_type == "lists":
        doc.add_paragraph('Bullet 1', style='List Bullet')
        doc.add_paragraph('Nested Bullet', style='List Bullet 2')
        doc.add_paragraph('Numbered 1', style='List Number')
        
    elif test_type == "mixed":
        doc.add_heading('Mixed Content', level=1)
        doc.add_paragraph('Some text.')
        doc.add_paragraph('A bullet', style='List Bullet')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Cell 00"
        table.cell(0, 1).text = "Cell 01"
        table.cell(1, 0).text = "Cell 10"
        table.cell(1, 1).text = "Cell 11"
        doc.add_paragraph('End text.')
        
    elif test_type == "header":
        section = doc.sections[0]
        header = section.header
        header.paragraphs[0].text = "CLASSIFIED — INTERNAL USE"
        doc.add_paragraph("Body text")

    elif test_type == "footer":
        section = doc.sections[0]
        footer = section.footer
        footer.paragraphs[0].text = "NTRO | Document ID: ABC-123"
        doc.add_paragraph("Body text")

    elif test_type == "header_footer":
        section = doc.sections[0]
        section.header.paragraphs[0].text = "Header text"
        section.footer.paragraphs[0].text = "Footer text"
        doc.add_paragraph("Body text")

    elif test_type == "multiple_sections":
        section = doc.sections[0]
        section.header.paragraphs[0].text = "Section 1 Header"
        doc.add_paragraph("Section 1 Body")
        
        new_section = doc.add_section()
        new_section.header.is_linked_to_previous = False
        new_section.header.paragraphs[0].text = "Section 2 Header"
        doc.add_paragraph("Section 2 Body")
        
    elif test_type == "shared_header":
        section = doc.sections[0]
        section.header.paragraphs[0].text = "Shared Header"
        doc.add_paragraph("Section 1 Body")
        
        new_section = doc.add_section()
        new_section.header.is_linked_to_previous = True
        doc.add_paragraph("Section 2 Body")
        
    elif test_type == "corrupted":
        return b"This is not a docx file."
        
    f = io.BytesIO()
    doc.save(f)
    return f.getvalue()

def test_docx_adapter_basic():
    file_bytes = create_synthetic_docx("basic")
    adapter = DOCXAdapter()
    cdo = adapter.parse(file_bytes, {"filename": "basic.docx"})
    
    assert cdo.modality == "docx"
    assert len(cdo.pages) == 1
    
    elements = cdo.pages[0].get("elements", [])
    # 1 section_start + 1 header + 1 footer + 3 body elements = 6
    assert len(elements) == 6
    
    # Find the heading element
    heading_element = next(e for e in elements if e["type"] == "heading")
    assert heading_element["text"] == "Title"
    
    # Find the paragraphs
    paragraphs = [e for e in elements if e["type"] == "paragraph"]
    assert len(paragraphs) >= 2
    assert paragraphs[0]["text"] == "Paragraph 1"

def test_docx_adapter_corrupted():
    file_bytes = create_synthetic_docx("corrupted")
    adapter = DOCXAdapter()
    with pytest.raises(CorruptedFileError):
        adapter.parse(file_bytes, {"filename": "bad.docx"})

def test_docx_pipeline_headings():
    file_bytes = create_synthetic_docx("headings")
    cdo = IngestionEngine.ingest_file(file_bytes, "headings.docx")
    
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    # We expect graph relationships to be preserved
    blocks = schema.semantic_blocks
    headings = [b for b in blocks if b.block_type == "heading"]
    paragraphs = [b for b in blocks if b.block_type == "paragraph"]
    
    assert len(headings) == 3
    assert len(paragraphs) == 3
    
    # Ensure paragraphs are CHILD_OF the correct headings
    rels = schema.relationships
    for p in paragraphs:
        # Each paragraph should point up to a parent (heading)
        child_edges = [r for r in rels if r.source_id == p.knowledge_object_id]
        assert len(child_edges) == 1
        assert child_edges[0].target_id in [h.knowledge_object_id for h in headings]

def test_docx_pipeline_tables():
    file_bytes = create_synthetic_docx("tables")
    cdo = IngestionEngine.ingest_file(file_bytes, "tables.docx")
    
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    blocks = schema.semantic_blocks
    tables = [b for b in blocks if b.block_type == "table"]
    rows = [b for b in blocks if b.block_type == "row"]
    cells = [b for b in blocks if b.block_type == "cell"]
    
    assert len(tables) == 1
    assert len(rows) == 3
    assert len(cells) == 12
    
    rels = schema.relationships
    
    # Cells point to Rows
    for cell in cells:
        edges = [r for r in rels if r.source_id == cell.knowledge_object_id]
        assert len(edges) == 1
        assert edges[0].target_id in [r.knowledge_object_id for r in rows]
        
    # Rows point to Table
    for row in rows:
        edges = [r for r in rels if r.source_id == row.knowledge_object_id]
        assert len(edges) == 1
        assert edges[0].target_id == tables[0].knowledge_object_id

def test_api_upload_docx():
    file_bytes = create_synthetic_docx("mixed")
    files_param = [('files', ('mixed.docx', file_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))]
    
    from app.core.security import create_access_token
    token = create_access_token({"sub": "admin_id", "username": "admin", "role": "SYSTEM_ADMIN"})
    headers = {"Authorization": f"Bearer {token}"}
    response = client.post("/api/upload", files=files_param, headers=headers)
    
    if response.status_code != 201:
        print("FAIL:", response.text)
    
    assert response.status_code == 201
    jobs = response.json()
    assert len(jobs) > 0
    
    job = jobs[0]
    assert job["filename"] == "mixed.docx"

def test_docx_adapter_header():
    file_bytes = create_synthetic_docx("header")
    cdo = IngestionEngine.ingest_file(file_bytes, "header.docx")
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    headers = [b for b in schema.semantic_blocks if b.block_type == "HEADER"]
    assert len(headers) == 1
    
    # Check if paragraph is extracted inside header
    header_id = headers[0].knowledge_object_id
    header_paragraphs = [b for b in schema.semantic_blocks if getattr(b, "parent_object_id", None) == header_id and b.block_type == "paragraph"]
    assert len(header_paragraphs) >= 1
    assert "CLASSIFIED" in header_paragraphs[0].content

def test_docx_adapter_footer():
    file_bytes = create_synthetic_docx("footer")
    cdo = IngestionEngine.ingest_file(file_bytes, "footer.docx")
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    footers = [b for b in schema.semantic_blocks if b.block_type == "FOOTER"]
    assert len(footers) == 1
    
    footer_id = footers[0].knowledge_object_id
    footer_paragraphs = [b for b in schema.semantic_blocks if getattr(b, "parent_object_id", None) == footer_id and b.block_type == "paragraph"]
    assert len(footer_paragraphs) >= 1
    assert "NTRO" in footer_paragraphs[0].content

def test_docx_adapter_shared_header():
    file_bytes = create_synthetic_docx("shared_header")
    cdo = IngestionEngine.ingest_file(file_bytes, "shared_header.docx")
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    headers = [b for b in schema.semantic_blocks if b.block_type == "HEADER"]
    # Even though there are two sections, the header is shared, so it should only be extracted once
    assert len(headers) == 1


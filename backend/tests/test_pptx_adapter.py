import pytest
import io
import pptx
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from fastapi.testclient import TestClient

from app.main import app
from app.core.ingestion.adapters.pptx_adapter import PPTXAdapter
from app.core.ingestion.engine import IngestionEngine
from app.core.pipeline.orchestrator import PipelineOrchestrator
from app.core.ingestion.detector import CorruptedFileError

client = TestClient(app)

def create_synthetic_pptx(test_type="basic") -> bytes:
    prs = pptx.Presentation()
    
    if test_type == "basic":
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Slide 1 Title"
        slide.placeholders[1].text = "Slide 1 Subtitle"
        
    elif test_type == "tables":
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Table Slide"
        x, y, cx, cy = Inches(2), Inches(2), Inches(4), Inches(1.5)
        table = slide.shapes.add_table(3, 3, x, y, cx, cy).table
        table.cell(0, 0).text = "Header 1"
        table.cell(1, 0).text = "Data 1"
        
    elif test_type == "charts":
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Chart Slide"
        
        chart_data = CategoryChartData()
        chart_data.categories = ['East', 'West', 'Midwest']
        chart_data.add_series('Series 1', (19.2, 21.4, 16.7))
        
        x, y, cx, cy = Inches(2), Inches(2), Inches(6), Inches(4.5)
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data)
        
    elif test_type == "notes":
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Has Notes"
        
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = "This is a speaker note."
        
    elif test_type == "mixed":
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Title"
        
        slide2 = prs.slides.add_slide(prs.slide_layouts[5])
        table = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(2), Inches(1)).table
        table.cell(0, 0).text = "Cell 1"
        
    elif test_type == "corrupted":
        return b"This is not a pptx file."
        
    f = io.BytesIO()
    prs.save(f)
    return f.getvalue()

def test_pptx_adapter_basic():
    file_bytes = create_synthetic_pptx("basic")
    adapter = PPTXAdapter()
    cdo = adapter.parse(file_bytes, {"filename": "basic.pptx"})
    
    assert cdo.modality == "pptx"
    assert len(cdo.pages) == 1
    
    elements = cdo.pages[0].get("elements", [])
    assert len(elements) == 2
    assert elements[0]["type"] == "title"
    assert elements[0]["text"] == "Slide 1 Title"

def test_pptx_adapter_corrupted():
    file_bytes = create_synthetic_pptx("corrupted")
    adapter = PPTXAdapter()
    with pytest.raises(CorruptedFileError):
        adapter.parse(file_bytes, {"filename": "bad.pptx"})

def test_pptx_pipeline_tables():
    file_bytes = create_synthetic_pptx("tables")
    cdo = IngestionEngine.ingest_file(file_bytes, "tables.pptx")
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    blocks = schema.semantic_blocks
    tables = [b for b in blocks if b.block_type == "table"]
    rows = [b for b in blocks if b.block_type == "row"]
    cells = [b for b in blocks if b.block_type == "cell"]
    
    assert len(tables) == 1
    assert len(rows) == 3
    assert len(cells) == 9
    
    rels = schema.relationships
    for cell in cells:
        edges = [r for r in rels if r.source_id == cell.knowledge_object_id]
        assert len(edges) == 1
        assert edges[0].target_id in [r.knowledge_object_id for r in rows]

def test_pptx_pipeline_charts():
    file_bytes = create_synthetic_pptx("charts")
    cdo = IngestionEngine.ingest_file(file_bytes, "charts.pptx")
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    blocks = schema.semantic_blocks
    charts = [b for b in blocks if b.block_type == "chart"]
    assert len(charts) == 1
    assert charts[0].metadata["chart_type"] == "COLUMN_CLUSTERED (51)"
    
    series = [b for b in blocks if b.block_type == "dataseries"]
    assert len(series) == 1
    assert "Series 1" in series[0].content
    
def test_pptx_pipeline_notes():
    file_bytes = create_synthetic_pptx("notes")
    cdo = IngestionEngine.ingest_file(file_bytes, "notes.pptx")
    schema, chunks = PipelineOrchestrator.process_document(cdo)
    
    blocks = schema.semantic_blocks
    notes = [b for b in blocks if b.metadata.get("is_speaker_notes")]
    assert len(notes) == 1
    assert "This is a speaker note." in notes[0].content

def test_api_upload_pptx():
    file_bytes = create_synthetic_pptx("mixed")
    files_param = [('files', ('mixed.pptx', file_bytes, 'application/vnd.openxmlformats-officedocument.presentationml.presentation'))]
    
    from app.core.security import create_access_token
    token = create_access_token({"sub": "admin_id", "username": "admin", "role": "SYSTEM_ADMIN"})
    headers = {"Authorization": f"Bearer {token}"}
    response = client.post("/api/upload", files=files_param, headers=headers)
    
    if response.status_code != 201:
        print("FAIL:", response.text)
    
    assert response.status_code == 201
    jobs = response.json()
    assert len(jobs) > 0
    assert jobs[0]["filename"] == "mixed.pptx"

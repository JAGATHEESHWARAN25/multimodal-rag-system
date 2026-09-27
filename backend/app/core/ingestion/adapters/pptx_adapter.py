import io
import logging
from typing import Dict, Any, Generator, Union, List
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from app.core.ingestion.adapters.base import BaseAdapter
from app.models.schema import CommonDocumentObject
from app.core.ingestion.detector import CorruptedFileError

logger = logging.getLogger(__name__)

class PPTXAdapter(BaseAdapter):
    """
    Adapter for processing Microsoft PowerPoint (.pptx) files.
    Extracts slides, shapes, text, tables, charts, and embedded images.
    Preserves bounding box spatial coordinates.
    """
    
    def __init__(self):
        super().__init__()
        self.modality = "pptx"
        
    def parse(self, file_bytes: bytes, file_metadata: Dict[str, Any]) -> Union[CommonDocumentObject, Generator[CommonDocumentObject, None, None]]:
        self.emit_event("AdapterParsingStarted", {"modality": self.modality, "filename": file_metadata.get("filename", "")})
        
        try:
            prs = pptx.Presentation(io.BytesIO(file_bytes))
        except PackageNotFoundError:
            raise CorruptedFileError("Invalid or corrupted PPTX file.")
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse PPTX: {str(e)}")
            
        pages = []
        
        for slide_idx, slide in enumerate(prs.slides):
            elements = []
            
            # Extract Slide Title if available
            if slide.shapes.title and slide.shapes.title.text:
                elements.append({
                    "type": "title",
                    "text": slide.shapes.title.text.strip(),
                    "bounding_box": self._get_bbox(slide.shapes.title)
                })
                
            # Process shapes
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue # Already processed
                    
                bbox = self._get_bbox(shape)
                
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        elements.append({
                            "type": "text_box",
                            "text": text,
                            "bounding_box": bbox,
                            "formatting": self._extract_formatting(shape)
                        })
                
                elif shape.has_table:
                    elements.append(self._parse_table(shape.table, bbox))
                    
                elif shape.has_chart:
                    elements.append(self._parse_chart(shape.chart, bbox))
                    
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image_bytes = shape.image.blob
                        mime_type = shape.image.content_type
                        elements.append({
                            "type": "image",
                            "image_bytes": image_bytes,
                            "mime_type": mime_type,
                            "bounding_box": bbox
                        })
                    except Exception as e:
                        logger.warning(f"Failed to extract image from slide {slide_idx}: {e}")
                        
                elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    elements.append({
                        "type": "group",
                        "text": "Grouped Shape",
                        "bounding_box": bbox
                    })
                
                else:
                    elements.append({
                        "type": "shape",
                        "shape_type": str(shape.shape_type),
                        "bounding_box": bbox
                    })
                    
            # Speaker Notes
            notes_text = ""
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    notes_text = notes_slide.notes_text_frame.text.strip()
                    if notes_text:
                        elements.append({
                            "type": "speaker_notes",
                            "text": notes_text
                        })
                        
            pages.append({
                "page_num": slide_idx,
                "elements": elements
            })
            
        import hashlib
        fallback_fp = hashlib.md5(file_bytes).hexdigest()
        
        cdo = CommonDocumentObject(
            document_id=file_metadata.get("document_id", "temp_doc_id"),
            fingerprint=file_metadata.get("fingerprint", fallback_fp),
            modality=self.modality,
            filename=file_metadata.get("filename", "unknown.pptx"),
            mime_type=file_metadata.get("mime_type", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
            metadata=file_metadata,
            pages=pages,
            text_blocks=[],
            tables=[]
        )
        
        self.emit_event("AdapterParsingCompleted", {"modality": self.modality})
        return cdo

    def _get_bbox(self, shape) -> Dict[str, Any]:
        """Extracts native EMU coordinates from a shape."""
        try:
            return {
                "left": shape.left,
                "top": shape.top,
                "width": shape.width,
                "height": shape.height
            }
        except:
            return None

    def _extract_formatting(self, shape) -> Dict[str, bool]:
        """Extracts basic semantic formatting from a text frame."""
        has_bold = False
        has_italic = False
        has_underline = False
        
        if not shape.has_text_frame:
            return {}
            
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.bold: has_bold = True
                if run.font.italic: has_italic = True
                if run.font.underline: has_underline = True
                
        return {
            "bold": has_bold,
            "italic": has_italic,
            "underline": has_underline
        }

    def _parse_table(self, table, bbox) -> Dict[str, Any]:
        """Parses a PPTX table into structured rows and cells."""
        rows_data = []
        for r_idx, row in enumerate(table.rows):
            cells_data = []
            for c_idx, cell in enumerate(row.cells):
                cells_data.append({
                    "row_index": r_idx,
                    "col_index": c_idx,
                    "text": cell.text_frame.text.strip() if cell.text_frame else ""
                })
            rows_data.append({
                "row_index": r_idx,
                "cells": cells_data
            })
            
        return {
            "type": "table",
            "bounding_box": bbox,
            "rows": rows_data,
            "num_rows": len(table.rows) if table.rows else 0,
            "num_cols": len(table.columns) if table.columns else 0
        }

    def _parse_chart(self, chart, bbox) -> Dict[str, Any]:
        """Parses a PPTX chart natively if possible."""
        series_data = []
        chart_title = ""
        if chart.has_title and chart.chart_title.has_text_frame:
            chart_title = chart.chart_title.text_frame.text
            
        try:
            for series in chart.series:
                series_data.append({
                    "name": series.name,
                    "values": list(series.values)
                })
        except Exception:
            pass # Complex charts might fail
            
        return {
            "type": "chart",
            "chart_type": str(chart.chart_type),
            "chart_title": chart_title,
            "bounding_box": bbox,
            "series": series_data
        }
